"""Attachment classification and mapping do not invoke AI or guess ambiguous roles."""
from pathlib import Path
from docx import Document
from fund_performance.attachments import (BENCHMARK, CONFIRM, DEMO_INPUT, FACTSHEET, KID, PEER, PRICE, SUPPORTING,
                                          add_files, apply_roles, classify, engine_inputs, rows, select_type,
                                          supporting_inputs, validate)

ROOT = Path(__file__).resolve().parents[1]


def test_price_is_read_and_ambiguous_factsheet_requires_confirmation(tmp_path):
    price = ROOT / 'demo/sample_index_fund/prices.csv'
    assert classify(str(price))[0] == PRICE
    doc = Document()
    doc.add_paragraph('Fund factsheet: Example fund')
    path = tmp_path / 'document.docx'
    doc.save(path)
    assert classify(str(path))[0] == CONFIRM
    kid = Document()
    kid.add_paragraph('Key Investor Information')
    kidpath = tmp_path / 'kid.docx'
    kid.save(kidpath)
    assert classify(str(kidpath))[0] == KID
    mixed = Document()
    mixed.add_paragraph('Fund factsheet. See the Key Investor Information document.')
    mixedpath = tmp_path / 'mixed.docx'
    mixed.save(mixedpath)
    assert classify(str(mixedpath))[0] == CONFIRM


def test_user_roles_removal_and_multiple_peers_map_exact_files(tmp_path):
    price = str(ROOT / 'demo/sample_index_fund/prices.csv')
    docs = []
    for name in ('fund', 'peer1', 'peer2'):
        p = tmp_path / f'{name}.docx'
        doc = Document()
        doc.add_paragraph('Fund factsheet and investment objectives')
        doc.save(p)
        docs.append(str(p))
    entries = add_files([], [price, *docs])
    table = rows(entries)
    for row, role in zip(table, [PRICE, FACTSHEET, PEER, PEER]):
        row[2] = role
    selected, error = apply_roles(entries, table)
    assert not error and not validate('Assess the fund', selected)
    assert engine_inputs(selected) == (price, docs[0], None, docs[1:], [])
    selected, error = apply_roles(entries, table[:-1])
    assert not error and engine_inputs(selected)[3] == [docs[1]]
    assert validate('', selected) == ['Enter a review question.']


def test_monthly_factsheet_reference_to_kid_is_not_classified_as_kid(tmp_path, monkeypatch):
    from fund_performance import attachments
    path = tmp_path / 'Monthly FactSheet.pdf'
    path.write_bytes(b'dummy')
    monkeypatch.setattr(attachments, 'read_document', lambda _: 'Monthly Fund Fact Sheet. See the KID for details.')
    assert classify(str(path))[0] == FACTSHEET
    monkeypatch.setattr(attachments, 'read_document', lambda _: 'Key Investor Information Document for the fund')
    assert classify(str(path))[0] == CONFIRM


def test_type_override_refreshes_validation_and_supports_multiple_documents(tmp_path, monkeypatch):
    from fund_performance import web
    monkeypatch.setattr(web, 'live_demo_allowed', lambda: True)
    files = []
    for name in ('fund', 'peer1', 'peer2', 'extra1', 'extra2'):
        path = tmp_path / f'{name}.docx'
        doc = Document()
        doc.add_paragraph('Fund factsheet and investment objectives')
        doc.save(path)
        files.append(str(path))
    price = str(ROOT / 'demo/sample_index_fund/prices.csv')
    entries = add_files([], [price, *files])
    table = rows(entries)
    status = {fn.name: fn.fn for fn in web.build_ui().fns.values()}['update_status']
    assert status('Review', entries, table, None)[1]['interactive'] is False
    for i, role in enumerate((FACTSHEET, PEER, PEER, SUPPORTING, SUPPORTING), 2):
        entries, table = select_type(entries, table, str(i), role)
    assert status('Review', entries, table, None)[1]['interactive'] is True
    assert engine_inputs(entries)[3] == files[1:3]
    assert supporting_inputs(entries) == files[3:]
    entries, table = select_type(entries, table, '2', KID)
    message, button = status('Review', entries, table, None)
    assert not button['interactive'] and 'main fund factsheet' in message


def test_dynamic_type_controls_create_private_dropdown_events(tmp_path):
    import gradio as gr
    from gradio.context import LocalContext
    from fund_performance.web import build_ui
    path = tmp_path / 'fund.docx'
    doc = Document()
    doc.add_paragraph('Fund factsheet and investment objectives')
    doc.save(path)
    entries = add_files([], [str(path)])
    ui = build_ui()
    renderable = ui.default_config.renderables[0]
    token = LocalContext.blocks_config.set(ui.default_config)
    try:
        renderable.apply(entries, rows(entries))
    finally:
        LocalContext.blocks_config.reset(token)
    controls = [b for b in ui.blocks.values() if isinstance(b, gr.Dropdown) and b.label == path.name]
    assert len(controls) == 1 and SUPPORTING in [value for _, value in controls[0].choices]
    assert FACTSHEET in [value for _, value in controls[0].choices]
    assert all(fn.api_visibility == 'private' for fn in ui.fns.values())


def test_unreadable_and_unsupported_files_are_blocked(tmp_path):
    bad = tmp_path / 'bad.pdf'
    bad.write_text('not a pdf')
    entries = add_files([], [str(bad)])
    assert entries[0]['role'] == CONFIRM
    rows_in = rows(entries)
    rows_in[0][2] = FACTSHEET
    selected, _ = apply_roles(entries, rows_in)
    assert any('could not be read' in issue for issue in validate('question', selected))
    rows_in[0][2] = BENCHMARK
    selected, _ = apply_roles(entries, rows_in)
    assert any('could not be read' in issue for issue in validate('question', selected))
    rows_in[0][2] = DEMO_INPUT
    assert apply_roles(entries, rows_in)[1]
    assert classify(str(tmp_path / 'unknown.txt'))[0] == CONFIRM
    benchmark = tmp_path / 'index.csv'
    benchmark.write_text('date,index_level\n2025-01-31,100\n2025-02-28,101\n')
    assert classify(str(benchmark))[0] == BENCHMARK


def test_synthetic_demo_callback_maps_prevalidated_inputs_and_rejects_edits():
    from fund_performance.web import build_ui
    ui = build_ui()
    callbacks = {fn.name: fn.fn for fn in ui.fns.values()}
    state, table, demo_id, mode, status = callbacks['choose_case']('sample_index_fund')
    assert demo_id == 'sample_index_fund' and mode == 'quick'
    message, button = callbacks['update_status']('What happened?', state, table, demo_id)
    assert button['interactive'] is True
    table = table[:-1]
    message, button = callbacks['update_status']('What happened?', state, table, demo_id)
    assert button['interactive'] is False and 'reload' in message


def test_unavailable_cases_never_substitute_synthetic_data(tmp_path, monkeypatch):
    from fund_performance import web
    monkeypatch.setenv('FUND_DEMO_DATA', str(tmp_path / 'absent'))
    callbacks = {fn.name: fn.fn for fn in web.build_ui().fns.values()}
    for demo_id in ('fidelity_us', 'vanguard_global'):
        state, table, chosen, mode, status = callbacks['choose_case'](demo_id)
        assert state == [] and table == [] and chosen is None and 'unavailable' in status


def test_synthetic_json_cannot_leak_into_a_custom_upload():
    from fund_performance.web import build_ui
    callbacks = {fn.name: fn.fn for fn in build_ui().fns.values()}
    state, table, demo_id, mode, _ = callbacks['choose_case']('sample_index_fund')
    message, button = callbacks['update_status']('Question?', state, table, None)
    assert button['interactive'] is False and 'only be used with the synthetic demo' in message


def test_question_cards_and_responsive_layout():
    from fund_performance.web import CSS, REVIEW_TYPES, build_ui
    assert list(REVIEW_TYPES) == ['overall', 'performance', 'fees_value', 'custom']
    assert 'repeat(4, minmax(0, 1fr))' in CSS
    assert '@media (max-width: 760px)' in CSS and 'grid-template-columns: 1fr' in CSS
    ui = build_ui()
    assert any(getattr(c, 'value', None) == 'overall' for c in ui.blocks.values())


def test_page_passes_edited_question_and_demo_inputs_to_existing_workflow(monkeypatch):
    import asyncio
    from fund_performance import web
    callbacks = {fn.name: fn.fn for fn in web.build_ui().fns.values()}
    state, table, demo_id, mode, _ = callbacks['choose_case']('sample_index_fund')
    seen = {}

    async def fake_review(question, price, factsheet, kid, peers, jsons, offline, limits, on_step):
        seen.update(question=question, price=price, jsons=jsons, offline=offline)
        for stage in range(1, 6):
            on_step(f'✓ Step {stage} of 5 · Stage {stage} · done')
        return web.Outcome('finished')

    monkeypatch.setattr(web, 'start_review', fake_review)
    async def collect():
        return [v async for v in callbacks['on_run']('My edited question?', state, table, demo_id, mode)]
    events = asyncio.run(collect())
    assert seen['question'] == 'My edited question?' and seen['offline'] is True
    assert len(seen['jsons']) == 3 and seen['price'].endswith('prices.csv')
    assert all(f'Step {n} of 5' in events[-1][0] for n in range(1, 6))


def test_dashboard_fallback_answer_and_research_sources_are_visible():
    from test_dashboard import data
    from fund_performance.dashboard import dashboard_html
    d = data()
    d['sources'] = [{'number': 1, 'query': 'Benchmark research', 'date': '2026-09-01', 'url': 'https://example.org'}]
    html = dashboard_html(d, 'analyst')
    assert 'Answer</p>' in html and 'The four ratings' in html
    assert 'Research sources' in html and '2026-09-01' in html


def test_page_requires_confirmed_identity_and_per_document_dates(tmp_path, monkeypatch):
    from fund_performance import web
    monkeypatch.setenv('OPENAI_API_KEY', 'test-key-not-real')
    monkeypatch.setenv('AI_RUNS_ENABLED', 'true')
    monkeypatch.delenv('PORT', raising=False)
    monkeypatch.delenv(web.OFFLINE_ENV, raising=False)
    from fund_performance.io import read_prices
    original = read_prices(ROOT / 'demo/sample_index_fund/prices.csv').prices
    price_file = tmp_path / 'prices.csv'
    original[['date', 'nav']].to_csv(price_file, index=False)
    price = str(price_file)
    doc = Document()
    doc.add_paragraph('Sample World benchmark evidence')
    docpath = tmp_path / 'benchmark.docx'
    doc.save(docpath)
    factsheet = tmp_path / 'fund.docx'
    doc.save(factsheet)
    series = tmp_path / 'benchmark.csv'
    original[['date', 'index_level']].to_csv(series, index=False)
    entries = add_files([], [price, str(factsheet), str(series), str(docpath)])
    table = rows(entries)
    for row, role in zip(table, [PRICE, FACTSHEET, BENCHMARK, BENCHMARK]):
        row[2] = role
    callbacks = {fn.name: fn.fn for fn in web.build_ui().fns.values()}
    message, button = callbacks['update_status']('Question?', entries, table, None)
    assert not button['interactive'] and 'name' in message and 'document date' in message
    table[-1][4] = '2026-09-01'
    message, button = callbacks['update_status']('Question?', entries, table, None, 'Sample World Index', 'GBP', True, 'market_close')
    assert button['interactive']


def test_page_passes_confirmed_benchmark_and_documents_to_review(tmp_path, monkeypatch):
    import asyncio
    from fund_performance import web
    from fund_performance.io import read_prices
    original = read_prices(ROOT / 'demo/sample_index_fund/prices.csv').prices
    price, series = tmp_path / 'prices.csv', tmp_path / 'index.csv'
    original[['date', 'nav']].to_csv(price, index=False)
    original[['date', 'index_level']].to_csv(series, index=False)
    factsheet, benchmark_doc = tmp_path / 'fund.docx', tmp_path / 'benchmark.docx'
    for path in (factsheet, benchmark_doc):
        doc = Document()
        doc.add_paragraph('Sample World Index factsheet')
        doc.save(path)
    entries = add_files([], [str(price), str(factsheet), str(series), str(benchmark_doc)])
    table = rows(entries)
    for row, role in zip(table, [PRICE, FACTSHEET, BENCHMARK, BENCHMARK]):
        row[2] = role
    table[-1][4] = '2026-09-01'
    callbacks = {fn.name: fn.fn for fn in web.build_ui().fns.values()}
    captured = {}

    async def fake_review(*args, benchmark_inputs=None):
        captured.update(benchmark_inputs)
        return web.Outcome('done')

    monkeypatch.setattr(web, 'start_review', fake_review)
    async def collect():
        return [v async for v in callbacks['on_run']('Question?', entries, table, None, web.LIVE,
            'Sample World Index', 'GBP', True, 'market_close')]
    import asyncio
    asyncio.run(collect())
    assert captured == {'series': str(series), 'documents': [{'path': str(benchmark_doc), 'date': '2026-09-01'}],
                        'name': 'Sample World Index', 'currency': 'GBP', 'confirmed': True,
                        'pricing_point': 'market_close'}


def test_selected_supporting_documents_reach_review(tmp_path, monkeypatch):
    import asyncio
    from fund_performance import web
    files = []
    for name in ('fund', 'extra1', 'extra2'):
        path = tmp_path / f'{name}.docx'
        doc = Document()
        doc.add_paragraph('Fund factsheet and investment objectives')
        doc.save(path)
        files.append(str(path))
    entries = add_files([], [str(ROOT / 'demo/sample_index_fund/prices.csv'), *files])
    table = rows(entries)
    for row, kind in zip(table, [PRICE, FACTSHEET, SUPPORTING, SUPPORTING]):
        row[2] = kind
    callbacks = {fn.name: fn.fn for fn in web.build_ui().fns.values()}
    captured = {}

    async def fake_review(*args, supporting_documents=None):
        captured['supporting'] = supporting_documents
        return web.Outcome('done')

    monkeypatch.setattr(web, 'start_review', fake_review)
    async def collect():
        return [v async for v in callbacks['on_run']('Review', entries, table, None, web.LIVE)]
    asyncio.run(collect())
    assert captured['supporting'] == files[1:]
